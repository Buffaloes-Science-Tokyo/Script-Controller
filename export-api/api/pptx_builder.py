"""Assembles held (fileId, slideIndex) picks from various source decks into
one exported .pptx.

python-pptx has no built-in "copy a slide from another presentation" -
`copy_slide` below is the standard workaround: deep-copy the source slide's
shape XML onto a blank layout in the destination, then re-create each
relationship the shapes reference (mainly embedded images) under a NEW rId
in the destination package and rewrite the copied XML's r:id/r:embed
attributes to point at it. Only image and external (hyperlink) relationships
are supported - a slide containing a chart, video, or embedded OLE object
raises NotImplementedError rather than silently producing a corrupt file.
"""
import copy
import io

from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT

import drive_download

R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"

_SKIPPED_RELTYPES = {RT.SLIDE_LAYOUT, RT.NOTES_SLIDE}
_UNSUPPORTED_RELTYPES = {RT.CHART, RT.OLE_OBJECT, RT.VIDEO, RT.MEDIA}


def _blank_layout(prs):
    for layout in prs.slide_layouts:
        if layout.name.strip().lower() == "blank":
            return layout
    return prs.slide_layouts[-1]


def _clear_placeholder_shapes(slide):
    for shape in list(slide.shapes):
        shape._element.getparent().remove(shape._element)


def _remap_rids(element, rid_map):
    for el in element.iter():
        for attr_name, attr_value in list(el.attrib.items()):
            if attr_name.startswith(f"{{{R_NS}}}") and attr_value in rid_map:
                el.set(attr_name, rid_map[attr_value])


def _copy_relationships(src_slide, dest_slide):
    """Re-creates each of src_slide's relationships in dest_slide's package
    and returns {old_rId: new_rId} for remapping the copied shape XML."""
    rid_map = {}
    for old_rId, rel in src_slide.part.rels.items():
        if rel.reltype in _SKIPPED_RELTYPES:
            continue
        if rel.reltype in _UNSUPPORTED_RELTYPES:
            raise NotImplementedError(
                f"copy_slide does not support relationship type {rel.reltype!r} "
                f"(slide contains a chart/video/OLE object it can't safely copy)"
            )
        if rel.is_external:
            new_rId = dest_slide.part.rels.get_or_add_ext_rel(rel.reltype, rel.target_ref)
        elif rel.reltype == RT.IMAGE:
            _, new_rId = dest_slide.part.get_or_add_image_part(io.BytesIO(rel.target_part.blob))
        else:
            raise NotImplementedError(
                f"copy_slide does not support relationship type {rel.reltype!r}"
            )
        rid_map[old_rId] = new_rId
    return rid_map


def copy_slide(src_prs, slide_index, dest_prs):
    """Deep-copies src_prs.slides[slide_index] into dest_prs, returning the new slide."""
    src_slide = src_prs.slides[slide_index]
    dest_slide = dest_prs.slides.add_slide(_blank_layout(dest_prs))
    _clear_placeholder_shapes(dest_slide)

    rid_map = _copy_relationships(src_slide, dest_slide)

    for shape in src_slide.shapes:
        shape_element = copy.deepcopy(shape._element)
        _remap_rids(shape_element, rid_map)
        dest_slide.shapes._spTree.insert_element_before(shape_element, "p:extLst")

    return dest_slide


def build_export(access_token, items):
    """items: ordered list of {"fileId": str, "slideIndex": int}.

    Downloads each distinct source file once (read-only, in-memory - never
    re-uploaded), copies the requested slides in order into a new
    presentation, and returns the assembled .pptx as bytes.
    """
    presentations_by_file_id = {}
    for item in items:
        file_id = item["fileId"]
        if file_id not in presentations_by_file_id:
            file_bytes = drive_download.download_file_bytes(access_token, file_id)
            presentations_by_file_id[file_id] = Presentation(io.BytesIO(file_bytes))

    dest_prs = Presentation()
    for item in items:
        src_prs = presentations_by_file_id[item["fileId"]]
        copy_slide(src_prs, item["slideIndex"], dest_prs)

    output = io.BytesIO()
    dest_prs.save(output)
    output.seek(0)
    return output.read()
