"""Standalone Flask app, deployed as its own Vercel project (Python runtime).

This service does exactly one thing: assemble held (fileId, slideIndex)
picks from Drive decks into a single exported .pptx. It's kept as Python
specifically to reuse the already-tested slide-copy logic (copy_slide),
which has no good equivalent library on the Node/Next.js side. Everything
else in the app (search, thumbnails, admin, the DB) lives in the Next.js
app - this service is deliberately as small as possible.

Everything lives in this one file (rather than split across sibling modules
under api/) because Vercel's Python runtime doesn't reliably bundle sibling
.py files alongside a function's entry file - `import pptx_builder` from a
neighboring api/pptx_builder.py raised ModuleNotFoundError at runtime even
though the build itself succeeded. This is the standard layout for a
Vercel Python function for that reason.

Deployed at its own Vercel domain; vercel.json rewrites every path here so
Flask's own router (not Vercel's) decides between /health and /export.

Called server-to-server only, from web/app/api/export/route.ts (the browser
never talks to this service directly and never sees the Google access
token), so there's no browser CORS concern to configure here.
"""
import copy
import io

from flask import Flask, jsonify, request, send_file
from google.oauth2.credentials import Credentials
from googleapiclient.discovery import build
from googleapiclient.http import MediaIoBaseDownload
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT

app = Flask(__name__)

R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"

_SKIPPED_RELTYPES = {RT.SLIDE_LAYOUT, RT.NOTES_SLIDE}
_UNSUPPORTED_RELTYPES = {RT.CHART, RT.OLE_OBJECT, RT.VIDEO, RT.MEDIA}


# ---------- Drive access ----------

def download_file_bytes(access_token, file_id):
    """Downloads raw bytes for a source .pptx, read-only, in-memory, never re-uploaded."""
    creds = Credentials(token=access_token)
    service = build("drive", "v3", credentials=creds, cache_discovery=False)

    media_request = service.files().get_media(fileId=file_id, supportsAllDrives=True)
    buffer = io.BytesIO()
    downloader = MediaIoBaseDownload(buffer, media_request)
    done = False
    while not done:
        _, done = downloader.next_chunk()
    buffer.seek(0)
    return buffer.read()


# ---------- pptx assembly ----------
#
# python-pptx has no built-in "copy a slide from another presentation" -
# copy_slide below is the standard workaround: deep-copy the source slide's
# shape XML onto a blank layout in the destination, then re-create each
# relationship the shapes reference (mainly embedded images) under a NEW rId
# in the destination package and rewrite the copied XML's r:id/r:embed
# attributes to point at it. Only image and external (hyperlink) relationships
# are supported - a slide containing a chart, video, or embedded OLE object
# raises NotImplementedError rather than silently producing a corrupt file.

def _blank_layout(prs):
    for layout in prs.slide_layouts:
        if layout.name.strip().lower() == "blank":
            return layout
    return prs.slide_layouts[-1]


def _clear_placeholder_shapes(slide):
    for shape in list(slide.shapes):
        shape._element.getparent().remove(shape._element)


def _remap_rids(element, rid_map):
    for el in element.iter():
        for attr_name, attr_value in list(el.attrib.items()):
            if attr_name.startswith(f"{{{R_NS}}}") and attr_value in rid_map:
                el.set(attr_name, rid_map[attr_value])


def _copy_relationships(src_slide, dest_slide):
    """Re-creates each of src_slide's relationships in dest_slide's package
    and returns {old_rId: new_rId} for remapping the copied shape XML."""
    rid_map = {}
    for old_rId, rel in src_slide.part.rels.items():
        if rel.reltype in _SKIPPED_RELTYPES:
            continue
        if rel.reltype in _UNSUPPORTED_RELTYPES:
            raise NotImplementedError(
                f"copy_slide does not support relationship type {rel.reltype!r} "
                f"(slide contains a chart/video/OLE object it can't safely copy)"
            )
        if rel.is_external:
            new_rId = dest_slide.part.rels.get_or_add_ext_rel(rel.reltype, rel.target_ref)
        elif rel.reltype == RT.IMAGE:
            _, new_rId = dest_slide.part.get_or_add_image_part(io.BytesIO(rel.target_part.blob))
        else:
            raise NotImplementedError(
                f"copy_slide does not support relationship type {rel.reltype!r}"
            )
        rid_map[old_rId] = new_rId
    return rid_map


def copy_slide(src_prs, slide_index, dest_prs):
    """Deep-copies src_prs.slides[slide_index] into dest_prs, returning the new slide."""
    src_slide = src_prs.slides[slide_index]
    dest_slide = dest_prs.slides.add_slide(_blank_layout(dest_prs))
    _clear_placeholder_shapes(dest_slide)

    rid_map = _copy_relationships(src_slide, dest_slide)

    for shape in src_slide.shapes:
        shape_element = copy.deepcopy(shape._element)
        _remap_rids(shape_element, rid_map)
        dest_slide.shapes._spTree.insert_element_before(shape_element, "p:extLst")

    return dest_slide


def build_export(access_token, items):
    """items: ordered list of {"fileId": str, "slideIndex": int}.

    Downloads each distinct source file once (read-only, in-memory - never
    re-uploaded), copies the requested slides in order into a new
    presentation, and returns the assembled .pptx as bytes.
    """
    presentations_by_file_id = {}
    for item in items:
        file_id = item["fileId"]
        if file_id not in presentations_by_file_id:
            file_bytes = download_file_bytes(access_token, file_id)
            presentations_by_file_id[file_id] = Presentation(io.BytesIO(file_bytes))

    dest_prs = Presentation()
    for item in items:
        src_prs = presentations_by_file_id[item["fileId"]]
        copy_slide(src_prs, item["slideIndex"], dest_prs)

    output = io.BytesIO()
    dest_prs.save(output)
    output.seek(0)
    return output.read()


# ---------- routes ----------

def _extract_access_token():
    header = request.headers.get("X-Google-Access-Token") or request.headers.get("Authorization", "")
    if header.startswith("Bearer "):
        header = header[len("Bearer "):]
    return header or None


@app.get("/health")
def health():
    return jsonify(status="ok")


@app.post("/export")
def export():
    access_token = _extract_access_token()
    if not access_token:
        return jsonify(error="missing X-Google-Access-Token header"), 401

    body = request.get_json(silent=True) or {}
    items = body.get("items") or []
    if not items:
        return jsonify(error="items must be a non-empty list of {fileId, slideIndex}"), 400

    try:
        pptx_bytes = build_export(access_token, items)
    except NotImplementedError as exc:
        return jsonify(error=str(exc)), 422

    return send_file(
        io.BytesIO(pptx_bytes),
        mimetype="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        as_attachment=True,
        download_name="script.pptx",
    )
