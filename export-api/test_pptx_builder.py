"""Unit tests for api/pptx_builder.copy_slide, isolated from Drive/the DB.

Builds small fixture presentations in-memory (no real files needed) covering
the two things copy_slide has to get right: shape XML deep-copy and
relationship remapping (embedded images), then asserts the copied slide in
the destination presentation matches the source.
"""
import io
import os
import sys

import pytest
from PIL import Image as PILImage
from pptx import Presentation
from pptx.util import Inches

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "api"))
import pptx_builder  # noqa: E402


def _make_png_bytes(color=(255, 0, 0)):
    buf = io.BytesIO()
    PILImage.new("RGB", (10, 10), color).save(buf, format="PNG")
    buf.seek(0)
    return buf


def _make_source_deck():
    prs = Presentation()
    blank = pptx_builder._blank_layout(prs)

    slide0 = prs.slides.add_slide(blank)
    slide0.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1)).text_frame.text = "slide 0"

    slide1 = prs.slides.add_slide(blank)
    textbox = slide1.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1))
    textbox.text_frame.text = "target slide"
    slide1.shapes.add_picture(_make_png_bytes(), Inches(2), Inches(2), Inches(1), Inches(1))

    return prs


def test_copy_slide_copies_text_and_image():
    src_prs = _make_source_deck()
    dest_prs = Presentation()

    copied = pptx_builder.copy_slide(src_prs, 1, dest_prs)

    assert len(dest_prs.slides) == 1
    texts = [s.text_frame.text for s in copied.shapes if s.has_text_frame]
    assert "target slide" in texts

    pictures = [s for s in copied.shapes if s.shape_type == 13]  # MSO_SHAPE_TYPE.PICTURE
    assert len(pictures) == 1
    # the image blob made it into the destination package as its own relationship
    assert pictures[0].image.blob == _make_png_bytes().getvalue()


def test_copy_slide_does_not_mutate_source():
    src_prs = _make_source_deck()
    dest_prs = Presentation()

    pptx_builder.copy_slide(src_prs, 0, dest_prs)

    assert len(src_prs.slides) == 2
    assert src_prs.slides[0].shapes[0].text_frame.text == "slide 0"


def test_copy_slide_multiple_picks_from_same_and_different_decks():
    deck_a = _make_source_deck()
    deck_b = _make_source_deck()
    dest_prs = Presentation()

    pptx_builder.copy_slide(deck_a, 0, dest_prs)
    pptx_builder.copy_slide(deck_a, 1, dest_prs)
    pptx_builder.copy_slide(deck_b, 1, dest_prs)

    assert len(dest_prs.slides) == 3
    assert dest_prs.slides[0].shapes[0].text_frame.text == "slide 0"
    assert "target slide" in [s.text_frame.text for s in dest_prs.slides[1].shapes if s.has_text_frame]
    assert "target slide" in [s.text_frame.text for s in dest_prs.slides[2].shapes if s.has_text_frame]

    output = io.BytesIO()
    dest_prs.save(output)
    output.seek(0)
    reloaded = Presentation(output)
    assert len(reloaded.slides) == 3


def test_copy_slide_raises_on_unsupported_relationship():
    prs = Presentation()
    blank = pptx_builder._blank_layout(prs)
    slide = prs.slides.add_slide(blank)
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE

    chart_data = CategoryChartData()
    chart_data.categories = ["a", "b"]
    chart_data.add_series("s1", (1, 2))
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1), Inches(4), Inches(3), chart_data)

    dest_prs = Presentation()
    with pytest.raises(NotImplementedError):
        pptx_builder.copy_slide(prs, 0, dest_prs)
